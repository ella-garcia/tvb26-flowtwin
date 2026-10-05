"""Read an uploaded CSV or XLSX into (headers, rows). Rows are {header: raw value}; blank rows are dropped."""
import csv
import io
from datetime import datetime


class ParseError(ValueError):
    pass


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _dedupe(headers):
    seen, out = {}, []
    for h in headers:
        h = (str(h).strip() if h is not None else "")
        n = seen.get(h, 0)
        seen[h] = n + 1
        out.append(h if n == 0 or not h else f"{h} ({n + 1})")
    return out


def parse_csv(data: bytes):
    text = _decode(data)
    first = next((l for l in text.splitlines() if l.strip()), "")
    delim = max(",;\t|", key=first.count)
    table = [r for r in csv.reader(io.StringIO(text), delimiter=delim) if any(c.strip() for c in r)]
    if not table:
        raise ParseError("The file is empty")
    return _to_rows(table)


def parse_xlsx(data: bytes):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ParseError(f"Not a readable .xlsx file: {e}") from e
    for ws in wb.worksheets:
        table = [list(r) for r in ws.iter_rows(values_only=True)
                 if any(c is not None and str(c).strip() != "" for c in r)]
        if table:
            return _to_rows(table)
    raise ParseError("The workbook has no data")


def _to_rows(table):
    headers = _dedupe(table[0])
    rows = []
    for i, r in enumerate(table[1:], start=2):
        rec = {h: (c.date() if isinstance(c, datetime) else c) for h, c in zip(headers, list(r) + [None] * (len(headers) - len(r)))}
        rec["__row__"] = i
        rows.append(rec)
    return headers, rows


def parse_file(data: bytes, file_name: str):
    n = (file_name or "").lower()
    if n.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        return parse_xlsx(data)
    if n.endswith(".xls"):
        raise ParseError("Old .xls files are not supported: save as .xlsx or .csv")
    return parse_csv(data)
